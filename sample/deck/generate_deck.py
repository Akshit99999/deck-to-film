#!/usr/bin/env python3
"""Generate a sample .pptx deck for smoke testing and demos.

Run: python generate_deck.py
Output: sample/deck/launchpad.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


def make_sample_deck() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    slides_data = [
        {
            "layout": 0,
            "title": "Launchpad",
            "body": "Ship products 10× faster",
            "notes": "Opening hook: Teams waste 80% of their time on infrastructure. Launchpad eliminates that.",
        },
        {
            "layout": 1,
            "title": "The Problem",
            "body": "Setup takes weeks\nDeployment is scary\nTeams are siloed\nAnalytics are expensive",
            "notes": "Focus on the pain: every startup re-solves the same infrastructure problem.",
        },
        {
            "layout": 1,
            "title": "Our Solution",
            "body": "One-click deploy\nBuilt-in collaboration\nPrivacy-first analytics\nZero config",
            "notes": "Each point maps to a pain point from the previous slide.",
        },
        {
            "layout": 1,
            "title": "Key Metrics",
            "body": "10× faster deployment\n3 min average setup time\n99.99% uptime SLA\n$0 to start",
            "notes": "These are aspirational benchmarks — only use real numbers you can defend.",
        },
        {
            "layout": 1,
            "title": "How It Works",
            "body": "Connect your repo → Launchpad detects your stack → One click to deploy → Live in seconds",
            "notes": "Walk through the 4-step user journey visually in the demo.",
        },
        {
            "layout": 0,
            "title": "Get Started Today",
            "body": "launchpad.example.com",
            "notes": "Call to action — free tier, no credit card required.",
        },
    ]

    for data in slides_data:
        layout = prs.slide_layouts[data["layout"]]
        slide = prs.slides.add_slide(layout)

        if slide.shapes.title:
            slide.shapes.title.text = data["title"]

        try:
            ph = slide.placeholders[1]
            ph.text = data["body"]
        except (KeyError, IndexError):
            pass

        notes = slide.notes_slide
        notes.notes_text_frame.text = data["notes"]

    out = Path(__file__).parent / "launchpad.pptx"
    prs.save(str(out))
    print(f"Sample deck written to {out}")
    return out


if __name__ == "__main__":
    make_sample_deck()
