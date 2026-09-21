"""Unit tests for PPTX parsing."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest

from videogen.ingest.pptx_parser import (
    ParsedDeck,
    ParsedSlide,
    SlideTable,
    parse_pptx,
)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


def _make_minimal_pptx(tmp_path: Path) -> Path:
    """Create a minimal valid .pptx file using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide_layout = prs.slide_layouts[0]

    # Slide 1: title + content
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Hello World"
    slide.placeholders[1].text = "This is body text\nSecond paragraph"

    # Add speaker notes
    notes = slide.notes_slide
    notes.notes_text_frame.text = "These are speaker notes for slide 1."

    # Slide 2: title only
    slide2 = prs.slides.add_slide(prs.slide_layouts[5])
    slide2.shapes.title.text = "Second Slide"

    pptx_path = tmp_path / "test.pptx"
    prs.save(str(pptx_path))
    return pptx_path


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


class TestParsePptx:
    def test_basic_parse(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        assert isinstance(deck, ParsedDeck)
        assert len(deck.slides) == 2
        assert deck.source_path == pptx_path

    def test_slide_titles_extracted(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        assert deck.slides[0].title == "Hello World"
        assert deck.slides[1].title == "Second Slide"

    def test_speaker_notes_extracted(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        assert "speaker notes" in deck.slides[0].speaker_notes.lower()

    def test_slide_indices_one_based(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        assert deck.slides[0].index == 1
        assert deck.slides[1].index == 2

    def test_full_text_includes_all_content(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        text = deck.full_text
        assert "Hello World" in text
        assert "Second Slide" in text

    def test_cache_key_is_stable(self, tmp_path: Path) -> None:
        pptx_path = _make_minimal_pptx(tmp_path)
        deck = parse_pptx(pptx_path)

        key1 = deck.cache_key
        key2 = deck.cache_key
        assert key1 == key2
        assert len(key1) == 16

    def test_missing_file_raises(self, tmp_path: Path) -> None:
        from pydantic import ValidationError

        with pytest.raises(Exception):  # PackageNotFoundError, FileNotFoundError, etc.
            parse_pptx(tmp_path / "nonexistent.pptx")

    def test_wrong_extension_raises(self, tmp_path: Path) -> None:
        bad_file = tmp_path / "deck.docx"
        bad_file.write_bytes(b"fake")
        with pytest.raises((ValueError, Exception)):
            # parse_pptx uses python-pptx which will fail on wrong format
            parse_pptx(bad_file)


class TestSlideTable:
    def test_table_str_repr(self) -> None:
        table = SlideTable(
            headers=["Name", "Score"],
            rows=[["Alice", "92"], ["Bob", "87"]],
        )
        assert table.headers == ["Name", "Score"]
        assert len(table.rows) == 2
